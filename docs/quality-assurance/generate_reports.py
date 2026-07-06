#!/usr/bin/env python3
"""Generate FRAIM quality report PPTX from markdown slide sources."""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def parse_slides(md_text: str) -> list[tuple[str, list[str]]]:
    """Parse ## Slide: headlines and bullet content."""
    slides: list[tuple[str, list[str]]] = []
    blocks = re.split(r"\n## Slide: ", md_text)
    for block in blocks[1:]:
        lines = block.strip().split("\n")
        headline = lines[0].strip()
        bullets = [ln.lstrip("- ").strip() for ln in lines[1:] if ln.strip().startswith("-")]
        slides.append((headline, bullets))
    return slides


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_content_slide(prs: Presentation, headline: str, bullets: list[str], notes: str) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = headline
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, bullet in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = notes


def build_pptx(md_path: Path, out_path: Path, report_title: str) -> None:
    md_text = md_path.read_text(encoding="utf-8")
    slides = parse_slides(md_text)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    today = date.today().isoformat()
    add_title_slide(
        prs,
        report_title,
        f"Camp Link Discovery Engine\nFRAIM Quality Assessment\n{today}",
    )
    for headline, bullets in slides:
        notes = f"Source: {md_path.name}\n{headline}"
        add_content_slide(prs, headline, bullets, notes)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"Wrote {out_path} ({len(slides) + 1} slides)")


def main() -> None:
    root = Path(__file__).resolve().parent.parent.parent
    qa = root / "docs" / "quality-assurance"
    today = date.today().isoformat()
    build_pptx(
        qa / f"code-quality-report-{today}.md",
        qa / f"code-quality-report-{today}.pptx",
        "Code Quality Assessment",
    )
    build_pptx(
        qa / f"test-quality-report-{today}.md",
        qa / f"test-quality-report-{today}.pptx",
        "Test Quality Assessment",
    )


if __name__ == "__main__":
    main()
